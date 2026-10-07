#!/usr/bin/env python3
"""Portable release checks. Does not open a browser or claim visual acceptance."""
import copy
import hashlib
import importlib.util
import io
import json
import re
import tempfile
import wave
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import unquote, urlsplit
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
BASE = DOCS / 'workpaper'
spec = importlib.util.spec_from_file_location('casework', ROOT/'case-fact-structuring/scripts/casework.py')
cw = importlib.util.module_from_spec(spec)
spec.loader.exec_module(cw)
checks = []
def check(name, condition):
    if not condition:
        raise AssertionError(name)
    checks.append(name)

class Markup(HTMLParser):
    def __init__(self, text):
        super().__init__()
        self.ids, self.links, self.active, self.resources = set(), [], [], []
        self.feed(text)
    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if attrs.get('id'): self.ids.add(attrs['id'])
        if attrs.get('href'): self.links.append(attrs['href'])
        if attrs.get('src'): self.links.append(attrs['src'])
        if attrs.get('aria-current') == 'page': self.active.append(attrs.get('href'))
        if tag == 'script' and attrs.get('src'): self.resources.append(attrs['src'])
        if tag == 'link' and attrs.get('rel') == 'stylesheet': self.resources.append(attrs.get('href',''))

pages = {p:Markup(p.read_text(encoding='utf-8')) for p in DOCS.glob('*.html')}
check('宣传、演示及使用页完整', {p.name for p in pages} == {'index.html','demo.html','guide.html'})
for p, markup in pages.items():
    check(p.name+' 当前位置标记', markup.active == [p.name])
    for href in markup.links:
        parsed = urlsplit(href)
        if parsed.scheme or parsed.netloc: continue
        target = (p.parent / unquote(parsed.path)).resolve() if parsed.path else p
        check(p.name+' 本地链接 '+href, target.is_relative_to(DOCS) and target.is_file())
        if parsed.fragment:
            if target in pages:
                check(p.name+' 锚点 '+parsed.fragment, parsed.fragment in pages[target].ids)
            elif target.name == '案件事实底稿.html':
                check('工作台栏目链接 '+parsed.fragment, parsed.fragment in {'overview','facts','materials','subjects','amounts','issues','changes'})
    text = p.read_text(encoding='utf-8')
    check(p.name+' 无外部脚本字体或嵌入底稿', '<iframe' not in text and all(not urlsplit(href).scheme for href in markup.resources))
    check(p.name+' 标示虚构及授权', '虚构' in text and '使用授权' in text)

data = json.loads((BASE/'case.json').read_text(encoding='utf-8'))
check('模拟案卷身份', data['case']['is_demo'] is True)
check('来源、类型、编号及金额计算校验', not cw.validate(data))
check('案卷规模与宣传一致', [len(data[k]) for k in ['materials','facts','citations','issues']] == [60,118,179,21])
check('材料读取范围一致', sum(m['state']=='已提取' for m in data['materials']) == 59)
html = (BASE/'案件事实底稿.html').read_text(encoding='utf-8')
embedded = json.loads(re.search(r'<script id="case-data" type="application/json">([\s\S]*?)</script>',html)[1])
web_snapshot=copy.deepcopy(data)
for material in web_snapshot['materials']:
    material.pop('path',None)
    for unit in material.get('units',[]): unit.pop('text',None)
check('网页事实引用及人工记录与JSON同源', embedded == web_snapshot)
check('协议流水主张分开', [a['amount'] for a in data['case']['amount_summaries'][:3]] == [68000000,56000000,26800000])
facts = {f['id']:f for f in data['facts']}
check('12条重点均为人工字段', sum(f.get('manual',{}).get('is_key') is True for f in data['facts'])==12)
check('人工分组覆盖6组', len({f.get('manual',{}).get('chronicle_group') for f in data['facts']} - {None,''}) == 6)
check('谈话未转换为确认欠款', '那50万后面统一处理' in facts['F0069']['description'] and '确认欠款' not in facts['F0069']['description'])
check('200万回款性质不同说法保留', '股东分配' in facts['F0061']['description'] and '扣减本金' in facts['F0071']['description'])
for m in data['materials']:
    src = BASE / m['path']
    check(m['id']+' 原件存在且哈希匹配', src.is_file() and hashlib.sha256(src.read_bytes()).hexdigest()==m['sha256'])
    check(m['id']+' 链接仍指原件', (BASE/unquote(m['source_href'])).resolve()==src.resolve())
audio_count = 0
for r in data['citations']:
    check(r['id']+' 引用原件可访问', (BASE/unquote(r['source_href'])).is_file())
    if r.get('audio_href'):
        with wave.open(str(BASE/unquote(r['audio_href']))) as audio:
            seconds = audio.getnframes()/audio.getframerate()
        check(r['id']+' 时间戳处于实际录音长度内', 0<=r['locator']['start']<r['locator']['end']<=seconds+.1)
        audio_count += 1

wb = load_workbook(BASE/'案件事实底稿.xlsx',data_only=True)
check('Excel保持七表', wb.sheetnames == ['案件速览','材料目录','主体表','事实大事记','原文与来源','冲突与待核','更新记录'])
cells = [c for sheet in wb for row in sheet for c in row if c.value is not None]
check('所有非空单元格使用华文仿宋', all(c.font.name == 'STFangsong' for c in cells))
check('大事记冻结编号和时间', wb['事实大事记'].freeze_panes=='C5')
sheet = wb['事实大事记']
heads = {c.value:c.column for c in sheet[4] if c.value}
check('人工分组与重点列位置正确', heads['复核状态']<heads['复核备注']<heads['整理分组']<heads['重点事项'])
rows = {sheet.cell(row,heads['事实编号']).value:row for row in range(5,sheet.max_row+1)}
check('Excel与JSON事实编号完整对应', set(rows)==set(facts))
for fid, row in rows.items():
    f = facts[fid]
    check(fid+' Excel摘要与金额对应', sheet.cell(row,heads['材料记载']).value==f['description'] and sheet.cell(row,heads['金额']).value==f.get('amount'))
    check(fid+' 复核及人工选择一致', sheet.cell(row,heads['复核状态']).value==f['review_status'] and (sheet.cell(row,heads['整理分组']).value or '')==f.get('manual',{}).get('chronicle_group','') and sheet.cell(row,heads['重点事项']).value==({True:'是',False:'否'}.get(f.get('manual',{}).get('is_key'))))
origins = wb['原文与来源']
check('Excel引用编号及连续原话对应', [(origins.cell(r,1).value,origins.cell(r,6).value) for r in range(5,origins.max_row+1)]==[(r['id'],r['quote']) for r in data['citations']])
for a in data['case']['amount_summaries']:
    check(a['id']+' Excel金额摘要一致', any(row[0]==a['label'] and row[1]==a['amount'] for row in wb['案件速览'].iter_rows(values_only=True)))
wb.close()

# Use disposable workbook copies: never modify the delivered Excel or case data.
with tempfile.TemporaryDirectory() as temp:
    edited = load_workbook(BASE/'案件事实底稿.xlsx')
    sheet = edited['事实大事记']
    row = rows['F0001']
    for title,value in [('人工更正金额',0),('复核备注','回导验收备注'),('整理分组','人工测试组'),('重点事项','否')]:
        sheet.cell(row,heads[title]).value = value
    review_file = Path(temp)/'reviews.xlsx'
    edited.save(review_file)
    reviewed=copy.deepcopy(data)
    cw.import_reviews(reviewed,review_file)
    fact = next(f for f in reviewed['facts'] if f['id']=='F0001')
    check('新字段及金额0回导', fact['manual']['amount']==0 and fact['manual']['chronicle_group']=='人工测试组' and fact['manual']['is_key'] is False and fact['review_note']=='回导验收备注')
    sheet.cell(row,heads['整理分组']).value = None
    edited.save(review_file)
    normal = copy.deepcopy(reviewed)
    cw.import_reviews(normal,review_file)
    check('普通空白不清除分组', normal['facts'][0]['manual']['chronicle_group']=='人工测试组')
    cleared = copy.deepcopy(reviewed)
    cw.import_reviews(cleared,review_file,clear_empty=True)
    check('显式清除空白分组', 'chronicle_group' not in cleared['facts'][0]['manual'] and cleared['facts'][0]['manual']['is_key'] is False)
    # Older column sets do not overwrite newly added selections.
    for title in sorted(['整理分组','重点事项'],key=lambda t:heads[t],reverse=True): sheet.delete_cols(heads[title])
    edited.save(review_file)
    legacy = copy.deepcopy(reviewed)
    cw.import_reviews(legacy,review_file)
    check('旧列集保留已有人工选择', legacy['facts'][0]['manual']['chronicle_group']=='人工测试组' and legacy['facts'][0]['manual']['is_key'] is False)
    previous = copy.deepcopy(reviewed)
    for round_index in range(3):
        candidate=copy.deepcopy(previous)
        candidate['facts'][0]['description'] = '新增版本建议 '+str(round_index)
        candidate['case']['revision'] += 1
        merged = cw.merge(previous,candidate)
        held = merged['facts'][0]
        check('第%d轮更新保护人工记录与原摘要'%(round_index+1), held['manual']['amount']==0 and held['manual']['is_key'] is False and held['manual']['chronicle_group']=='人工测试组' and held['review_note']=='回导验收备注' and held['description']==facts['F0001']['description'] and bool(merged.get('pending_changes')))
        previous=merged
    edited.close()

source = ROOT/'case-fact-structuring'
source_files = [p for p in source.rglob('*') if p.is_file() and '__pycache__' not in p.parts]
for platform in ['codex','workbuddy']:
    with zipfile.ZipFile(DOCS/'downloads'/f'{platform}-lugou-v0.3.zip') as package:
        check(platform+' 安装包包含全部源文件', all('case-fact-structuring/'+str(p.relative_to(source)) in package.namelist() for p in source_files))
        check(platform+' 安装包与当前源文件相同', all(package.read('case-fact-structuring/'+str(p.relative_to(source)))==p.read_bytes() for p in source_files))

# Inspect public plaintext and compressed XML/text/PDF contents, including nested ZIPs.
bad_paths = re.compile(rb'/' + rb'Users/[^/\s]+/|' + rb'/' + rb'private/var/folders/')
secrets = re.compile(rb'gh[pousr]_[A-Za-z0-9]{20,}|sk-proj-[A-Za-z0-9_-]{20,}')
denied = {'runtime.local.json','.runtime','node_modules','_support','__pycache__','.env','.DS_Store'}
scanned=0
def privacy_bytes(name, raw, depth=0):
    global scanned
    check('公开文件无私有运行配置 '+name, not denied.intersection(Path(name).parts))
    check('公开文件无本机路径或密钥 '+name, not bad_paths.search(raw) and not secrets.search(raw))
    scanned += 1
    if zipfile.is_zipfile(io.BytesIO(raw)):
        check('嵌套压缩深度有限 '+name, depth<4)
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            for member in z.namelist():
                if member.endswith('/'): continue
                check('压缩条目路径安全 '+name, not member.startswith('/') and '..' not in Path(member).parts)
                privacy_bytes(name+'!'+member,z.read(member),depth+1)
    elif name.lower().endswith('.pdf'):
        try:
            reader=PdfReader(io.BytesIO(raw))
            text = str(reader.metadata or '') + '\n'.join(page.extract_text() or '' for page in reader.pages)
            check('PDF文字及元数据无本机路径 '+name, not bad_paths.search(text.encode()) and not secrets.search(text.encode()))
        except Exception:
            # The mock unreadable PDF intentionally contains no valid page structure.
            check('只允许演示中的损坏PDF读取失败', '损坏' in name or '不可读' in name)

for p in ROOT.rglob('*'):
    rel=p.relative_to(ROOT)
    if not p.is_file() or any(part in {'.git','__pycache__'} for part in rel.parts) or p.name.endswith('.local.json'): continue
    privacy_bytes(str(rel),p.read_bytes())

report={'checks':len(checks),'passed':True,'scope':'公开发布静态、数据、来源、Excel与人工记录保护检查；不包含浏览器视觉或实际听核',
        'counts':{'materials':60,'facts':118,'citations':179,'issues':21,'excel_populated_cells':len(cells),'audio_citations':audio_count,'scanned_entries':scanned},
        'privacy':'公开源包、页面、原件、Excel和全部压缩包扫描通过；无本机绝对路径、运行配置或已识别密钥格式'}
(ROOT/'tests/public-results.local.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
