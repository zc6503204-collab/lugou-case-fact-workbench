#!/usr/bin/env python3
"""Local material extraction, provenance, review merge, and rendering."""
import argparse
import copy
import datetime as dt
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import uuid
import zipfile
from email import policy
from email.parser import BytesParser
import xml.etree.ElementTree as ET
from urllib.parse import quote
from zoneinfo import ZoneInfo

SKILL = Path(__file__).resolve().parents[1]
STATUSES = {"未复核", "已核对原件", "需补核", "存在异议"}
KINDS = {"材料记载", "一方陈述", "待核推断"}
ISSUES = {"冲突", "缺口", "识别疑点", "待确认变更"}
AMOUNT_ROLES = {"协议总额","协议分项","流水发生额","回款记载","一方主张","主张拆项","财务口径","其他金额"}

def now():
    return dt.datetime.now(ZoneInfo("Asia/Shanghai")).isoformat(timespec="seconds")

def read_json(path):
    p=Path(path).resolve()
    data=json.loads(p.read_text(encoding="utf-8"))
    if isinstance(data,dict) and data.get("schema_version")=="1.0":
        for m in data.get("materials",[]):
            if m.get("path") and not Path(m["path"]).is_absolute():
                m["path"]=str((p.parent/m["path"]).resolve())
    return data

def save_json(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    temp.replace(path)

def config():
    p = Path(os.environ.get("CASEWORK_CONFIG", SKILL / "runtime.local.json"))
    c = read_json(p) if p.exists() else {}
    base = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies"
    for k, v in {"python": base/"python/bin/python3", "node":base/"node/bin/node",
                 "node_modules":base/"node/node_modules"}.items():
        if not c.get(k) and v.exists():
            c[k] = str(v)
    if not c.get("ffmpeg"):
        v = shutil.which("ffmpeg")
        fallback = Path("/Applications/Downie 4.app/Contents/Resources/ffmpeg")
        if v or fallback.exists():
            c["ffmpeg"] = v or str(fallback)
    return c

def next_id(items, prefix):
    nums = [int(x["id"][1:]) for x in items if re.fullmatch(prefix + r"\d+", x.get("id",""))]
    return prefix + str(max(nums, default=0)+1).zfill(4)

def empty_case(name, demo=False):
    return {"schema_version":"1.0", "case":{"name":name,"is_demo":demo,
            "case_id":str(uuid.uuid4()),"revision":1,
            "scope_note":"事实整理底稿；材料表述与案件认定分别处理。", "updated_at":now()},
            "materials":[], "subjects":[], "facts":[], "citations":[], "issues":[], "changes":[]}

def unit(mid, key, text, locator, verification="文字提取，待核原件"):
    return {"id":mid+":"+key, "text":text, "locator":locator, "verification":verification,
            "reading":{"status":"未读","note":""}}

def ocr(path, page=None):
    env = os.environ.copy()
    env["CLANG_MODULE_CACHE_PATH"] = str(Path(tempfile.gettempdir())/"casework-swift-cache")
    args = ["/usr/bin/swift", str(SKILL/"scripts/macos_ocr.swift"), str(path)]
    if page is not None:
        args += [str(page)]
    r = subprocess.run(args, env=env, capture_output=True, text=True, timeout=180)
    if r.returncode:
        raise RuntimeError("OCR 不可用：" + r.stderr[-800:])
    return json.loads(r.stdout)

def transcribe(path, out, language="zh"):
    c = config()
    for k in ("asr_python", "asr_model", "ffmpeg"):
        if not c.get(k) or not Path(c[k]).exists():
            raise RuntimeError("本地录音转写缺少配置或文件：" + k)
    env = os.environ.copy()
    env.update(HF_HUB_OFFLINE="1", TRANSFORMERS_OFFLINE="1",
               TOKENIZERS_PARALLELISM="false", NUMBA_CACHE_DIR=str(Path(out)/"_numba"))
    args = [c["asr_python"], str(SKILL/"scripts/transcribe.py"), str(path),
            "--out",str(out),"--model",c["asr_model"],"--ffmpeg",c["ffmpeg"],"--language",language]
    r = subprocess.run(args, env=env, capture_output=True, text=True, timeout=7200)
    if r.returncode:
        raise RuntimeError("转写失败：" + r.stderr[-1200:])
    return read_json(Path(out)/"transcript.json")

def extract(path, mid, out, only_pages=None):
    suffix = path.suffix.lower()
    units, notes, method = [], [], ""
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        method = "PDF 文字层"
        for i, page in enumerate(reader.pages, 1):
            if only_pages is not None and i not in only_pages: continue
            text = (page.extract_text() or "").strip()
            verification = "文字提取，待核原件"
            loc = {"kind":"pdf","page":i,"printed_page":None}
            # Sparse text layers and images can conceal a scanned body behind a page number.
            if len(normalized(text)) < 60 or bool(page.images):
                try:
                    result = ocr(path, i)
                    recognized = result["text"].strip()
                    if recognized:
                        text = recognized
                    verification = "自动识别，待核原件"
                    loc["ocr_blocks"] = result["blocks"]
                    method = "PDF 文字层/OCR"
                except Exception as e:
                    notes.append("实际第%d页：%s" % (i,e))
                    verification = "文字层保留，OCR失败，需补核"
            if text:
                units.append(unit(mid,"p"+str(i),text,loc,verification))
            else:
                notes.append("实际第%d页未取得文字" % i)
        return units,notes,method,len(reader.pages)
    if suffix == ".docx":
        from docx import Document
        doc = Document(path)
        method = "Word 段落/表格"
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        pi,ti = 0,0
        for child in doc.element.body:
            if child.tag.endswith('}p'):
                pi += 1
                text=Paragraph(child,doc).text
                if text.strip(): units.append(unit(mid,"para"+str(pi),text,{"kind":"word","paragraph":pi}))
            elif child.tag.endswith('}tbl'):
                ti += 1
                for ri,row in enumerate(Table(child,doc).rows,1):
                    text="\n".join("第%d列：%s"%(ci,c.text) for ci,c in enumerate(row.cells,1))
                    units.append(unit(mid,"t%d-r%d"%(ti,ri),text,{"kind":"word","table":ti,"row":ri}))
        with zipfile.ZipFile(path) as z:
            xml=z.read('word/document.xml').decode('utf-8')
            for tag,label in [('txbxContent','文本框'),('drawing','嵌入图片'),('ins','修订插入'),('del','修订删除')]:
                if re.search(r'<w:'+tag+r'[ >]',xml): notes.append(label+'须查看 Word 原件，未当作普通正文完整读取')
            for n in z.namelist():
                if re.fullmatch(r'word/(header\d+|footer\d+|footnotes|endnotes)\.xml',n):
                    tree=ET.fromstring(z.read(n))
                    txt='\n'.join(x.text for x in tree.iter() if x.tag.endswith('}t') and x.text)
                    if txt: units.append(unit(mid,n.replace('/','-'),txt,{"kind":"word","part":n}))
            if '<w:tbl' in xml and any(c.tables for t in doc.tables for r in t.rows for c in r.cells):
                notes.append('嵌套表格须查看原件，顶层行文本可能未覆盖嵌套内容')
        return units,notes,method,None
    if suffix in {".xlsx",".xlsm"}:
        from openpyxl import load_workbook
        wb = load_workbook(path,data_only=False,read_only=True)
        cached = load_workbook(path,data_only=True,read_only=True)
        method = "Excel 工作表/单元格"
        with zipfile.ZipFile(path) as z:
            if any(n.startswith('xl/media/') for n in z.namelist()): notes.append('工作簿含嵌图，须核对图像内容')
        for si,ws in enumerate(wb.worksheets,1):
            for row in ws.iter_rows():
                entries = []
                for cell in row:
                    if cell.value is None:
                        continue
                    val = cell.value
                    if cell.data_type == "f":
                        cv = cached[ws.title][cell.coordinate].value
                        val = "公式 %s；缓存 %s"%(val,cv if cv is not None else "缺失，需核算")
                    if isinstance(val,(dt.datetime,dt.date)):
                        val = val.isoformat()
                    entries.append("%s：%s"%(cell.coordinate,val))
                if entries:
                    rnum = row[0].row
                    units.append(unit(mid,"s%d-r%d"%(si,rnum),"\n".join(entries),
                        {"kind":"excel","sheet":ws.title,"range":row[0].coordinate+":"+row[-1].coordinate,
                         "sheet_state":ws.sheet_state,"formats":{c.coordinate:c.number_format for c in row if c.value is not None}}))
        wb.close()
        cached.close()
        return units,notes,method,None
    if suffix == '.eml':
        msg=BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        for key in ('From','To','Date','Subject'):
            if msg.get(key): units.append(unit(mid,'header-'+key,key+': '+str(msg[key]),{"kind":"email","part":key}))
        for i,part in enumerate(msg.walk(),1):
            if part.get_content_disposition()=='attachment':
                notes.append('邮件附件未独立提交：'+str(part.get_filename() or '未命名附件'))
            elif part.get_content_type()=='text/plain':
                for j,line in enumerate(part.get_content().splitlines(),1):
                    if line.strip(): units.append(unit(mid,'body%d-l%d'%(i,j),line,{"kind":"email","part":"正文%d"%i,"line":j}))
            elif part.get_content_type()=='text/html': notes.append('邮件 HTML 正文须核对原件；优先提取纯文本部分')
        return units,notes,'邮件 MIME 解码',None
    if suffix in {".txt",".md",".csv",".tsv",".srt",".vtt"}:
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("gb18030")
        method = "逐行文本"
        for i,line in enumerate(text.splitlines(),1):
            if line.strip():
                units.append(unit(mid,"l"+str(i),line,{"kind":"text","line":i}))
        return units,notes,method,None
    if suffix in {".png",".jpg",".jpeg",".tif",".tiff",".heic",".bmp"}:
        if suffix in {'.tif','.tiff'}:
            from PIL import Image
            with Image.open(path) as im:
                if getattr(im,'n_frames',1)>1: notes.append('多页 TIFF 仅识别首帧，后续帧须补充为独立图像或 PDF')
        result = ocr(path)
        for i,b in enumerate(result["blocks"],1):
            units.append(unit(mid,"ocr"+str(i),b["text"],{"kind":"image","block":i,
                "box":b["box"]},"自动识别，待核原件"))
        return units,notes,"macOS Vision OCR",None
    if suffix in {".wav",".mp3",".m4a",".aac",".aiff",".aif",".flac",".ogg",".mp4"}:
        result = transcribe(path,Path(out)/"_extraction"/mid)
        for i,s in enumerate(result["segments"],1):
            units.append(unit(mid,"seg"+str(i),s["text"],{"kind":"audio","start":s["start"],
                "end":s["end"],"duration":result["duration"],"speaker":"未确认"},"自动转写，待听核"))
        if not units:
            notes.append("未检出可转写语音；原录音仍需核听")
        return units,notes,"本地 MLX Whisper",None
    raise RuntimeError("暂不支持此格式；请提供可读取版本")

def collect(args):
    resume=getattr(args,'resume',False) and Path(args.out).exists()
    data = read_json(args.out) if resume else read_json(args.previous) if args.previous else empty_case(args.name or "案件事实底稿",args.demo)
    if not data['case'].get('case_id'): data['case']['case_id']=str(uuid.uuid4())
    data['case']['revision']=data['case'].get('revision',1)+(1 if args.previous and not resume else 0)
    if args.name:
        data["case"]["name"] = args.name
    mats = data["materials"]
    inputs = []
    excluded=[]
    missing_paths=[]
    outpath=Path(args.out).resolve()
    output_dir=outpath.parent
    for item in args.input:
        p = Path(item).expanduser().resolve()
        if not p.exists():
            missing_paths.append(str(p))
            mid = next_id(mats,"M")
            mats.append({"id":mid,"filename":p.name,"path":str(p),"sha256":None,
                "state":"原文件缺失","method":"","notes":["指定路径不存在"],"units":[],
                "duplicate_of":None,"version_of":None})
            continue
        if p.is_dir():
            for x in sorted(p.rglob('*')):
                if not x.is_file(): continue
                rel=x.relative_to(p)
                reason=None
                if any(z.startswith('.') for z in rel.parts): reason='隐藏路径'
                elif x.resolve()==outpath or (output_dir!=p and output_dir.is_relative_to(p) and x.is_relative_to(output_dir)): reason='成果输出目录'
                elif rel.parts[0] in {'_support','_extraction','sources','media'} or x.name in {'案件事实底稿.html','案件事实底稿.xlsx'}: reason='生成成果或中间文件'
                if reason: excluded.append({'path':str(x),'reason':reason})
                else: inputs.append(x)
        else:
            inputs.append(p)
    for p in dict.fromkeys(inputs):
        try: digest = hashlib.sha256(p.read_bytes()).hexdigest()
        except OSError as e:
            mats.append({'id':next_id(mats,'M'),'filename':p.name,'path':str(p),'sha256':None,'state':'无法读取','method':'','notes':[str(e)],'units':[]})
            save_json(args.out,data)
            continue
        exact = next((m for m in mats if m["path"]==str(p) and m.get("sha256")==digest),None)
        if exact:
            if args.retry_unreadable and (not exact.get("units") or exact.get('state') in {'部分提取','无法读取','未识别'}):
                try:
                    retry_pages={int(x) for note in exact.get('notes',[]) for x in re.findall(r'实际第(\d+)页',note)}
                    us,notes,method,pages=extract(p,exact["id"],Path(args.out).parent,
                        only_pages=retry_pages if p.suffix.lower()=='.pdf' and retry_pages else None)
                    old={u['id']:u for u in exact.get('units',[])}
                    for u in us:
                        if u['id'] not in old: old[u['id']]=u
                        elif normalized(old[u['id']]['text'])!=normalized(u['text']):
                            u['id'] += ':retry'+str(len(old)); old[u['id']]=u
                            notes.append('重试文字与旧提取不同，旧原话保留；新单元须回原件确认')
                    exact.update(units=list(old.values()),notes=notes,method=method,page_count=pages,
                                 state="部分提取" if us and notes else "已提取" if us else "未识别")
                except Exception as e:
                    exact.update(state="无法读取",notes=[str(e)])
                data["changes"].append({"at":now(),"type":"重试识别","target_id":exact["id"],"detail":exact["state"]})
                save_json(args.out,data)
            continue
        prior = next((m for m in reversed(mats) if m["path"]==str(p)),None)
        duplicate = next((m for m in mats if m.get("sha256")==digest),None)
        mid = next_id(mats,"M")
        m = {"id":mid,"filename":p.name,"path":str(p),"sha256":digest,"state":"已提取",
             "method":"","notes":[],"duplicate_of":duplicate["id"] if duplicate else None,
             "version_of":prior["id"] if prior else None,"units":[],"page_count":None}
        if prior and prior.get("sha256") != digest:
            prior["state"] = "已被新版本替换，旧原件待提供"
            prior.setdefault("notes",[]).append("当前路径已变更为 "+mid)
        try:
            us,notes,method,pages = extract(p,mid,Path(args.out).parent)
            m.update(units=us,notes=notes,method=method,page_count=pages)
            if notes:
                m["state"] = "部分提取" if us else "未识别"
        except Exception as e:
            m.update(state="无法读取",notes=[str(e)])
        mats.append(m)
        if prior:
            affected=[f["id"] for f in data["facts"] if any(c["material_id"]==prior["id"] and c["id"] in f.get("citation_ids",[]) for c in data["citations"])]
            data["changes"].append({"at":now(),"type":"旧来源失效","target_id":prior["id"],"detail":"当前路径已被新内容覆盖，旧引用保留；旧原件需补提供。"})
            if affected:
                data["issues"].append({"id":next_id(data["issues"],"Q"),"kind":"待确认变更",
                    "title":"来源出现新版本，关联事实需回查","detail":prior["id"]+" → "+mid+"；旧原件待提供，未自动改写已有事实。",
                    "fact_ids":affected,"material_ids":[prior["id"],mid],"citation_ids":[],
                    "next_step":"对照新旧版本逐项核对；保留人工复核记录。","review_status":"未复核","review_note":""})
        data["changes"].append({"at":now(),"type":"新增版本" if prior else "新增材料",
            "target_id":mid,"detail":p.name+("；重复 "+duplicate["id"] if duplicate else "")})
        save_json(args.out,data)
    for m in mats:
        if not Path(m["path"]).exists() and m["state"] != "原文件缺失":
            m["state"] = "原文件缺失"
            m.setdefault("notes",[]).append("本次未能定位原文件；保留旧提取")
    data["case"]["updated_at"] = now()
    for m in mats:
        m['extraction']={'status':m['state'],'notes':m.get('notes',[]),
                         'obtained_units':len(m.get('units',[])),'page_count':m.get('page_count')}
    data['collection']={'input_paths':[str(Path(p).expanduser().resolve()) for p in args.input],
        'files_seen':len(set(inputs))+len(excluded),'included_files':len(set(inputs)),
        'excluded':excluded,'missing_paths':missing_paths,'completed':True,'at':now()}
    save_json(args.out,data)
    print(json.dumps({"output":str(Path(args.out).resolve()),"materials":len(mats),
                      "readable":sum(bool(m["units"]) for m in mats)},ensure_ascii=False))

def normalized(s):
    return re.sub(r"\s+","",str(s))

def validate(data):
    errors = []
    for key in ("materials","subjects","facts","citations","issues","changes"):
        if not isinstance(data.get(key),list):
            errors.append("缺少或格式错误："+key+" 必须为数组")
    if not isinstance(data.get("case"),dict):errors.append("case 必须为对象")
    if errors:return errors
    if data.get("schema_version") != "1.0":
        errors.append("schema_version 必须为 1.0")
    maps = {}
    for key,prefix in (("materials","M"),("subjects","S"),("facts","F"),("citations","R"),("issues","Q")):
        items = data.get(key,[])
        ids = [x.get("id") for x in items]
        if len(set(ids)) != len(ids):
            errors.append(key+" 有重复编号")
        if any(not re.fullmatch(prefix+r"\d{4,}",str(i)) for i in ids):
            errors.append(key+" 编号格式错误")
        maps[key] = {x.get("id"):x for x in items}
    for key in ("case","changes"):
        if key not in data:
            errors.append("缺少 "+key)
    def refs(values,key,owner):
        for v in values:
            if v not in maps[key]:
                errors.append(owner+" 引用了不存在的 "+str(v))
    for m in maps["materials"].values():
        for key in ('category','batch_id','document_form'):
            if key in m and not isinstance(m[key],str): errors.append(m['id']+' '+key+' 应为文字')
        uids = [u.get("id") for u in m.get("units",[])]
        if len(set(uids))!=len(uids):
            errors.append(m["id"]+" 提取单元重复")
        for rel in ("version_of","duplicate_of"):
            if m.get(rel):
                refs([m[rel]],"materials",m["id"])
                if m[rel]==m["id"]:
                    errors.append(m["id"]+" 不得引用自身为版本或重复")
        for u in m.get('units',[]):
            if u.get('reading',{}).get('status','未读') not in {'未读','已读'}: errors.append(u['id']+' 阅读状态错误')
    for s in maps["subjects"].values():
        refs(s.get("material_ids",[]),"materials",s["id"])
    for c in maps["citations"].values():
        refs([c.get("material_id")],"materials",c["id"])
        m = maps["materials"].get(c.get("material_id"),{})
        u = next((u for u in m.get("units",[]) if u["id"]==c.get("unit_id")),None)
        if not u:
            errors.append(c["id"]+" 缺少真实提取单元")
            continue
        if not normalized(c.get("quote","")) or normalized(c["quote"]) not in normalized(u["text"]):
            errors.append(c["id"]+" 原话不在对应提取单元中")
        loc = u.get("locator",{})
        if loc.get("kind")=="audio" and not (0 <= loc.get("start",-1) < loc.get("end",-1) <= loc.get("duration",0)+0.1):
            errors.append(c["id"]+" 录音时间戳越界")
    for f in maps["facts"].values():
        if f.get('amount_role') and f['amount_role'] not in AMOUNT_ROLES: errors.append(f['id']+' 金额角色错误')
        for key in ('agreement_round','transaction_id'):
            if key in f and not isinstance(f[key],str): errors.append(f['id']+' '+key+' 应为文字')
        refs(f.get('checked_citation_ids',[]),'citations',f['id'])
        if not f.get("citation_ids"):
            errors.append(f["id"]+" 没有原文引用")
        refs(f.get("citation_ids",[]),"citations",f["id"])
        refs(f.get("subject_ids",[]),"subjects",f["id"])
        if f.get("kind") not in KINDS:
            errors.append(f["id"]+" 事实类型错误")
        if f.get("review_status","未复核") not in STATUSES:
            errors.append(f["id"]+" 复核状态错误")
        a = f.get("amount")
        if a is not None and (isinstance(a,bool) or not isinstance(a,(int,float))):
            errors.append(f["id"]+" 金额应为数值")
        date = f.get("date",{})
        precision,iso = date.get("precision"),date.get("iso")
        formats = {"day":r"\d{4}-\d{2}-\d{2}","month":r"\d{4}-\d{2}","year":r"\d{4}"}
        if precision in formats:
            if not isinstance(iso,str) or not re.fullmatch(formats[precision],iso):
                errors.append(f["id"]+" 日期精度与值不符")
            else:
                try:
                    if precision=="day": dt.date.fromisoformat(iso)
                    elif precision=="month": dt.date.fromisoformat(iso+"-01")
                    else: dt.date(int(iso),1,1)
                except ValueError:
                    errors.append(f["id"]+" 日期值无效")
        elif precision not in {"unknown","range"} or iso is not None:
            errors.append(f["id"]+" 未知/范围日期不能填精确日期")
        manual = f.get("manual",{})
        if not isinstance(manual,dict):
            errors.append(f['id']+' manual 应为对象')
            continue
        if 'chronicle_group' in manual and not isinstance(manual['chronicle_group'],str):
            errors.append(f['id']+' 整理分组应为文字')
        if 'is_key' in manual and not isinstance(manual['is_key'],bool):
            errors.append(f['id']+' 重点事项应为布尔值')
        if manual.get("amount") is not None and (isinstance(manual["amount"],bool) or not isinstance(manual["amount"],(int,float))):
            errors.append(f["id"]+" 人工金额应为数值")
    for q in maps["issues"].values():
        if q.get("kind") not in ISSUES:
            errors.append(q["id"]+" 问题类型错误")
        for field,key in (("fact_ids","facts"),("material_ids","materials"),("citation_ids","citations")):
            refs(q.get(field,[]),key,q["id"])
        refs(q.get('received_material_ids',[]),'materials',q['id'])
        if q.get('progress','待补核') not in {'待补核','已收到待核','已核实关闭'}: errors.append(q['id']+' 办理状态错误')
        if q.get('progress')=='已核实关闭' and not q.get('resolution_note'): errors.append(q['id']+' 关闭问题须记录核查结果')
        if q.get('progress')=='已核实关闭' and q.get('review_status')!='已核对原件': errors.append(q['id']+' 关闭问题须先核查关联原件')
    for a in data['case'].get('amount_summaries',[]):
        if isinstance(a.get('amount'),bool) or not isinstance(a.get('amount'),(int,float)): errors.append('金额摘要须为数值')
        if not a.get('citation_ids'): errors.append('金额摘要 '+a.get('id','')+' 缺少引用')
        refs(a.get('citation_ids',[]),'citations',a.get('id','摘要'))
        refs(a.get('fact_ids',[]),'facts',a.get('id','摘要'))
        terms=a.get('calculation',{}).get('terms',[])
        if terms:
            refs([t.get('fact_id') for t in terms],'facts',a.get('id','摘要'))
            if any(t.get('sign') not in {1,-1} for t in terms): errors.append('摘要计算符号错误')
            if len({t.get('fact_id') for t in terms})!=len(terms): errors.append('摘要计算重复事实')
            txs=[maps['facts'].get(t.get('fact_id'),{}).get('transaction_id') for t in terms]
            known=[x for x in txs if x]
            if len(set(known))!=len(known): errors.append('摘要计算重复交易')
            try:
                total=sum(maps['facts'][t['fact_id']]['amount']*t['sign'] for t in terms)
                if abs(total-a['amount'])>0.01: errors.append('金额摘要 '+a.get('id','')+' 与拆项计算不一致')
            except (KeyError,TypeError): errors.append('摘要计算缺少有效金额')
    return errors

def import_reviews(data, path, clear_empty=False, allow_legacy=False):
    from openpyxl import load_workbook
    wb = load_workbook(path,data_only=True,read_only=True)
    original=data
    data=copy.deepcopy(data)
    meta={}
    if '案件速览' in wb.sheetnames:
        for row in wb['案件速览'].iter_rows(min_row=1,max_row=60,values_only=True):
            if len(row)>1 and row[0] in {'案件标识','底稿版本'}: meta[row[0]]=row[1]
    if data['case'].get('case_id'):
        if not meta.get('案件标识'):
            if not allow_legacy: raise ValueError('Excel 缺少案件标识；请使用对应版本底稿。旧版迁移须显式 --allow-legacy-reviews。')
        elif str(meta['案件标识'])!=data['case']['case_id']: raise ValueError('Excel 属于另一案件，停止回导')
        elif str(meta.get('底稿版本'))!=str(data['case'].get('revision',1)): raise ValueError('Excel 版本与待更新底稿不同；请以该 Excel 对应的 case.json 为上一版')
    mapping = {"事实大事记":("facts","事实编号"),"冲突与待核":("issues","问题编号")}
    imported = 0
    for sheet,(key,idhead) in mapping.items():
        if sheet not in wb.sheetnames:
            raise ValueError('Excel 缺少 '+sheet)
        rows = list(wb[sheet].iter_rows(values_only=True))
        headeridx = next((i for i,r in enumerate(rows[:10]) if idhead in r),None)
        if headeridx is None:
            raise ValueError(sheet+' 缺少编号表头')
        heads = {h:i for i,h in enumerate(rows[headeridx]) if h is not None}
        if not {'复核状态','复核备注'}.issubset(heads): raise ValueError(sheet+' 缺少指定复核列')
        records = {x["id"]:x for x in data[key]}
        seen=set()
        for row in rows[headeridx+1:]:
            rid = row[heads[idhead]]
            if not rid: continue
            if rid in seen: raise ValueError(sheet+' 重复编号 '+str(rid))
            seen.add(rid)
            if rid not in records:
                raise ValueError(sheet+' 未知编号 '+str(rid))
            item = records[rid]
            touched = False
            def cell(name):
                idx = heads.get(name)
                return row[idx] if idx is not None and idx<len(row) else None
            s = cell("复核状态")
            if s:
                if s not in STATUSES:
                    raise ValueError(str(rid)+" 复核状态不在允许值中")
                item["review_status"] = s
                touched = True
            note = cell("复核备注")
            if note is not None or clear_empty:
                item["review_note"] = str(note) if note is not None else ""
                touched = True
            if key=="facts":
                manual = item.setdefault("manual",{})
                for title,field in (("人工更正事实","description"),("人工更正时间","date_raw"),
                                    ("人工更正金额","amount"),("人工更正币种","currency")):
                    v = cell(title)
                    if v is not None or clear_empty:
                        if field=="amount" and v is not None:
                            if isinstance(v,bool) or not isinstance(v,(float,int)):
                                raise ValueError(str(rid)+" 人工更正金额请填写数值")
                        manual[field] = v
                        touched = True
                # Missing columns in older workbooks never clear new human selections.
                if '整理分组' in heads:
                    group=cell('整理分组')
                    if group is not None:
                        if not isinstance(group,str): raise ValueError(str(rid)+' 整理分组请填写文字')
                        if group.strip() or clear_empty:
                            manual['chronicle_group']=group.strip();touched=True
                    elif clear_empty:
                        manual.pop('chronicle_group',None);touched=True
                if '重点事项' in heads:
                    flag=cell('重点事项')
                    if flag is not None and flag!='':
                        if isinstance(flag,bool): value=flag
                        elif isinstance(flag,str) and flag.strip() in {'是','否'}: value=flag.strip()=='是'
                        else: raise ValueError(str(rid)+' 重点事项请填写“是”或“否”')
                        manual['is_key']=value;touched=True
                    elif clear_empty:
                        manual.pop('is_key',None);touched=True
                checked=cell('核查引用')
                if checked is not None:
                    item['checked_citation_ids']=[x for x in re.split(r'[、,，\s]+',str(checked)) if x]
                    touched=True
            for title,field in [('复核人','reviewed_by'),('复核时间','reviewed_at')]:
                v=cell(title)
                if v is not None: item[field]=v.isoformat() if hasattr(v,'isoformat') else str(v); touched=True
            if key=='issues':
                for title,field in [('办理状态','progress'),('核查结果','resolution_note')]:
                    v=cell(title)
                    if v is not None: item[field]=str(v); touched=True
                received=cell('收补材料')
                if received is not None: item['received_material_ids']=[x for x in re.split(r'[、,，\s]+',str(received)) if x];touched=True
            if touched:
                imported += 1
    wb.close()
    errors=validate(data)
    if errors: raise ValueError('\n'.join(errors))
    data["changes"].append({"at":now(),"type":"导入复核","target_id":"","detail":"导入 %d 条复核记录"%imported})
    original.clear();original.update(data)
    return imported

def protected(f):
    return f.get("review_status","未复核")!="未复核" or bool(f.get("review_note")) or any(
        v is not None and v!="" for v in f.get("manual",{}).values())

def merge(previous,candidate,reviews=None):
    if previous['case'].get('case_id') and candidate['case'].get('case_id')!=previous['case']['case_id']:
        raise ValueError('候选数据属于另一案件，停止合并')
    data = copy.deepcopy(previous)
    if reviews:
        import_reviews(data,reviews)
    data["case"] = copy.deepcopy(candidate["case"])
    for key in ("materials","subjects","citations","issues"):
        records = {x["id"]:x for x in data[key]}
        for new in candidate[key]:
            old = records.get(new["id"])
            if key=="citations" and old and old!=new:
                raise ValueError("引用编号不可改指其它原话："+new["id"]+"；请新增 R 编号")
            if key=="materials" and old and old.get("sha256")!=new.get("sha256"):
                raise ValueError("材料编号不可覆盖新内容："+new["id"]+"；请新增 M 编号")
            merged = copy.deepcopy(new)
            if key=="issues" and old:
                for field in ("review_status","review_note","reviewed_by","reviewed_at","resolution_note"):
                    if field in old: merged[field]=old[field]
                if old.get('progress')=='已核实关闭' or old.get('review_status','未复核')!='未复核' or old.get('review_note'):
                    merged['progress']=old.get('progress','待补核')
            records[new["id"]] = merged
        data[key] = list(records.values())
    records = {f["id"]:f for f in data["facts"]}
    for new in candidate["facts"]:
        old = records.get(new["id"])
        if old:
            fields = ("date","subject_ids","description","amount","currency","kind","citation_ids",
                      "agreement_round","transaction_id","amount_role")
            changed = any(old.get(k)!=new.get(k) for k in fields)
            if changed and protected(old):
                if any(p.get("fact_id")==old["id"] and p.get("proposed")==new for p in data.get("pending_changes",[])):
                    continue
                data.setdefault("pending_changes",[]).append({"fact_id":old["id"],"proposed":copy.deepcopy(new),"at":now()})
                qid = next_id(data["issues"],"Q")
                def proposed_text(f):
                    amount="未列金额" if f.get("amount") is None else str(f["amount"])+" "+f.get("currency","")
                    return f.get("date",{}).get("raw","")+"；"+f.get("description","")+"；"+amount
                data["issues"].append({"id":qid,"kind":"待确认变更","title":"含复核或人工记录的事实出现新内容",
                    "detail":old["id"]+" 保留原记录及人工更正。\n原记录："+proposed_text(old)+"\n新建议："+proposed_text(new),"fact_ids":[old["id"]],
                    "material_ids":[],"citation_ids":new.get("citation_ids",[]),
                    "next_step":"回查新引用及原复核依据，确认采纳哪一版；人工复核与更正继续保留。",
                    "review_status":"未复核","review_note":""})
                data["changes"].append({"at":now(),"type":"待确认变更","target_id":old["id"],"detail":qid})
                continue
            result = copy.deepcopy(new)
            for k in ("review_status","review_note","manual","reviewed_by","reviewed_at","checked_citation_ids"):
                if k in old: result[k] = copy.deepcopy(old[k])
            records[new["id"]] = result
            if changed:
                data["changes"].append({"at":now(),"type":"更新事实","target_id":new["id"],"detail":"未复核记录按本次材料更新"})
        else:
            records[new["id"]] = copy.deepcopy(new)
            data["changes"].append({"at":now(),"type":"新增事实","target_id":new["id"],"detail":new["description"]})
    data["facts"] = list(records.values())
    held={p['fact_id'] for p in data.get('pending_changes',[]) if not p.get('decision')}
    old_summaries={a.get('id'):a for a in previous['case'].get('amount_summaries',[])}
    for i,a in enumerate(data['case'].get('amount_summaries',[])):
        if held.intersection(a.get('fact_ids',[])) and a.get('id') in old_summaries:
            data['case']['amount_summaries'][i]=copy.deepcopy(old_summaries[a['id']])
    if 'collection' in candidate: data['collection']=copy.deepcopy(candidate['collection'])
    # Preserve new collection audit entries without duplicating history.
    for change in candidate.get("changes",[]):
        if change not in data["changes"]: data["changes"].append(change)
    data["case"]["updated_at"] = now()
    errs = validate(data)
    if errs:
        raise ValueError("\n".join(errs))
    return data

def decide_change(data,fact_id,decision,note):
    pending=[p for p in data.get('pending_changes',[]) if p['fact_id']==fact_id and not p.get('decision')]
    if len(pending)!=1: raise ValueError('须有且仅有一个未处理建议；多版本建议请先逐一核查并整理')
    p=pending[0]
    if decision=='accept':
        old=next(f for f in data['facts'] if f['id']==fact_id)
        p['before']=copy.deepcopy(old)
        kept={k:copy.deepcopy(old[k]) for k in ('manual','review_note','reviewed_by','reviewed_at','checked_citation_ids') if k in old}
        old.clear();old.update(copy.deepcopy(p['proposed']));old.update(kept)
        old['review_status']='需补核'
    p.update(decision=decision,decision_note=note,decided_at=now())
    data['changes'].append({'at':now(),'type':'采纳变更' if decision=='accept' else '拒绝变更','target_id':fact_id,'detail':note})
    for q in data['issues']:
        if q['kind']=='待确认变更' and q.get('fact_ids')==[fact_id]:
            q.update(progress='已核实关闭',resolution_note=note or decision,review_status='已核对原件')
    return data

def locator_text(loc):
    kind = loc.get("kind")
    if kind=="pdf":
        s = "PDF 实际第%d页"%loc["page"]
        return s+("；文内第%s页"%loc["printed_page"] if loc.get("printed_page") is not None else "；文内页码未标定")
    if kind=="word":
        if loc.get('part'): return 'Word '+loc['part']
        return "第%d段"%loc["paragraph"] if "paragraph" in loc else "表%d第%d行"%(loc["table"],loc["row"])
    if kind=="excel": return "%s!%s"%(loc["sheet"],loc["range"])
    if kind=="text": return "第%d行"%loc["line"]
    if kind=='email': return '邮件 '+loc.get('part','')+(' 第%d行'%loc['line'] if loc.get('line') else '')
    if kind=="image": return "图像 OCR 块%d"%loc["block"]
    if kind=="audio": return "%.1f–%.1f 秒；说话人%s"%(loc["start"],loc["end"],loc.get("speaker","未确认"))
    return "定位待核"

def enrich(data,out,bundle=False):
    d = copy.deepcopy(data)
    if bundle and d.get('collection'):
        d['collection']['input_paths']=['原提交路径已移除；sources 包含本包累计原件']
        for e in d['collection'].get('excluded',[]): e['path']=Path(e['path']).name
        d['collection']['missing_paths']=[Path(p).name for p in d['collection'].get('missing_paths',[])]
    mats = {m["id"]:m for m in d["materials"]}
    for m in mats.values():
        p = Path(m["path"])
        exists = p.exists()
        same = exists and (not m.get("sha256") or hashlib.sha256(p.read_bytes()).hexdigest()==m["sha256"])
        if not same:
            m["source_href"] = None
            m["source_note"] = "原件缺失或当前路径内容已变更"
            if bundle:m["path"]="missing/"+p.name
        elif bundle:
            target = Path(out)/"sources"/(m["id"]+"_"+m['filename'])
            target.parent.mkdir(parents=True,exist_ok=True)
            if p.resolve()!=target.resolve(): shutil.copy2(p,target)
            m["source_href"] = "sources/"+quote(target.name)
            # Public fictional examples should not embed the author's home path.
            m["path"] = "sources/"+target.name
        else:
            m["source_href"] = quote(os.path.relpath(p,out),safe="/")
        if same and any(u.get("locator",{}).get("kind")=="audio" for u in m.get("units",[])):
            # A browser-compatible listening copy, with the original time axis.
            ffmpeg = config().get("ffmpeg")
            if ffmpeg:
                audio = Path(out)/"media"/(m["id"]+".wav")
                audio.parent.mkdir(parents=True,exist_ok=True)
                subprocess.run([ffmpeg,"-nostdin","-v","error","-y","-i",str(p),
                    "-ar","16000","-ac","1","-c:a","pcm_s16le",str(audio)],
                    check=True,capture_output=True)
                m["audio_href"] = "media/"+audio.name
    for c in d["citations"]:
        m = mats[c["material_id"]]
        u = next(u for u in m["units"] if u["id"]==c["unit_id"])
        c.update(locator=u["locator"],location=locator_text(u["locator"]),
                 verification=u["verification"],filename=m["filename"],source_href=m.get("source_href"),
                 audio_href=m.get("audio_href"))
    return d

def render(data,args):
    errs = validate(data)
    if errs:
        raise ValueError("\n".join(errs))
    out = Path(args.out).resolve()
    out.mkdir(parents=True,exist_ok=True)
    save_json(out/"case.json",data)
    enriched = enrich(data,out,args.bundle_sources)
    if args.bundle_sources: save_json(out/'case.json',enriched)
    # HTML carries selected quotes and coverage states, not hidden full extraction text.
    web=copy.deepcopy(enriched)
    for m in web['materials']:
        m.pop('path',None)
        for u in m.get('units',[]):
            u.pop('text',None)
            u.get('locator',{}).pop('ocr_blocks',None)
    for r in web['citations']: r.get('locator',{}).pop('ocr_blocks',None)
    payload = json.dumps(web,ensure_ascii=False).replace("<","\\u003c").replace("&","\\u0026")
    template = (SKILL/"assets/workpaper.html").read_text(encoding="utf-8")
    (out/"案件事实底稿.html").write_text(template.replace("__CASE_DATA__",payload),encoding="utf-8")
    support = out/"_support"
    support.mkdir(exist_ok=True)
    save_json(support/"render-data.json",enriched)
    c = config()
    if not c.get("node") or not c.get("node_modules"):
        raise RuntimeError("网页已生成；Excel 渲染缺少 node/node_modules 配置")
    modules = support/"node_modules"
    if not modules.exists(): modules.symlink_to(c["node_modules"],target_is_directory=True)
    shutil.copy2(SKILL/"scripts/build_xlsx.mjs",support/"build_xlsx.mjs")
    r = subprocess.run([c["node"],str(support/"build_xlsx.mjs"),str(support/"render-data.json"),str(out)],
                       capture_output=True,text=True,timeout=240)
    if r.returncode:
        raise RuntimeError("Excel 渲染失败："+r.stderr[-2000:])
    subprocess.run([sys.executable,str(SKILL/"scripts/add_xlsx_links.py"),str(out/"案件事实底稿.xlsx"),
                    str(support/"xlsx-links.json")],check=True,capture_output=True,text=True)
    print(r.stdout.strip())

def main():
    c = config()
    if c.get("python") and Path(sys.executable).resolve()!=Path(c["python"]).resolve() and not os.environ.get("CASEWORK_NO_REEXEC"):
        os.execv(c["python"],[c["python"],str(Path(__file__).resolve())]+sys.argv[1:])
    p = argparse.ArgumentParser(description="律构·案件事实梳理")
    sub = p.add_subparsers(dest="command",required=True)
    sub.add_parser("doctor")
    a=sub.add_parser('audit');a.add_argument('case')
    a = sub.add_parser("collect")
    a.add_argument("--input",action="append",required=True)
    a.add_argument("--out",required=True)
    a.add_argument("--name")
    a.add_argument("--previous")
    a.add_argument("--demo",action="store_true")
    a.add_argument("--retry-unreadable",action="store_true",help="重试没有提取单元的材料，保留材料编号")
    a.add_argument('--resume',action='store_true',help='从本次输出断点恢复')
    a = sub.add_parser("validate"); a.add_argument("case")
    a = sub.add_parser("render"); a.add_argument("case"); a.add_argument("--out",required=True)
    a.add_argument("--bundle-sources",action="store_true")
    a = sub.add_parser("transcribe"); a.add_argument("input"); a.add_argument("--out",required=True)
    a.add_argument("--language",default="zh")
    a = sub.add_parser("merge")
    for k in ("previous","candidate","out"): a.add_argument("--"+k,required=True)
    a.add_argument("--reviews")
    a = sub.add_parser("import-reviews")
    a.add_argument("case"); a.add_argument("--reviews",required=True); a.add_argument("--out",required=True)
    a.add_argument("--clear-empty",action="store_true")
    a.add_argument('--allow-legacy-reviews',action='store_true',help='仅用于用户已明确指定的旧版 Excel 迁移')
    a=sub.add_parser('decide-change')
    a.add_argument('case');a.add_argument('--fact-id',required=True);a.add_argument('--decision',choices=['accept','reject'],required=True)
    a.add_argument('--note',required=True);a.add_argument('--out',required=True)
    args = p.parse_args()
    if args.command=="doctor":
        print(json.dumps({"config":c,"libraries":{k:bool(importlib.util.find_spec(k)) for k in
            ("pypdf","docx","openpyxl")},"swift":Path("/usr/bin/swift").exists(),
            "asr_configured":all(c.get(k) and Path(c[k]).exists() for k in ("asr_python","asr_model","ffmpeg"))},
            ensure_ascii=False,indent=2))
    elif args.command=="collect": collect(args)
    elif args.command=='audit':
        data=read_json(args.case);units=[u for m in data['materials'] for u in m.get('units',[])]
        print(json.dumps({'case':data['case']['name'],'errors':validate(data),
            'collection_completed':data.get('collection',{}).get('completed',False),
            'material_gaps':[{'id':m['id'],'state':m['state'],'notes':m.get('notes',[])} for m in data['materials'] if m['state']!='已提取'],
            'units':len(units),'unread_units':[u['id'] for u in units if u.get('reading',{}).get('status')!='已读'],
            'facts_unreviewed':[f['id'] for f in data['facts'] if f.get('review_status')!='已核对原件'],
            'open_issues':[q['id'] for q in data['issues'] if q.get('progress')!='已核实关闭'],
            'note':'程序校验、模型阅读与律师原件复核分别记录，不代表法院事实认定。'},ensure_ascii=False,indent=2))
    elif args.command=="validate":
        errs = validate(read_json(args.case))
        print(json.dumps({"valid":not errs,"errors":errs},ensure_ascii=False,indent=2))
        if errs: sys.exit(1)
    elif args.command=="render": render(read_json(args.case),args)
    elif args.command=="transcribe": print(json.dumps(transcribe(Path(args.input),args.out,args.language),ensure_ascii=False))
    elif args.command=="merge": save_json(args.out,merge(read_json(args.previous),read_json(args.candidate),args.reviews))
    elif args.command=="import-reviews":
        data = read_json(args.case)
        import_reviews(data,args.reviews,args.clear_empty,args.allow_legacy_reviews)
        errs = validate(data)
        if errs: raise ValueError("\n".join(errs))
        # Standalone import becomes a new snapshot; merge sets its own version.
        data['case']['revision']=int(data['case'].get('revision') or 1)+1
        data['case']['updated_at']=now()
        save_json(args.out,data)
    elif args.command=='decide-change':
        data=decide_change(read_json(args.case),args.fact_id,args.decision,args.note)
        errs=validate(data)
        if errs: raise ValueError('\n'.join(errs))
        save_json(args.out,data)

if __name__=="__main__":
    try:
        main()
    except Exception as e:
        print(str(e),file=sys.stderr)
        sys.exit(1)
